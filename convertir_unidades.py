"""Reconstruye una capa de concentraciones en mM sin alterar fuentes ni
imputar."""

import argparse
import json
import sqlite3
import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from pypdf import PdfReader

from apoyo import huella

RAIZ = Path(__file__).resolve().parent


def main():
    argumentos = argparse.ArgumentParser(description=__doc__)
    argumentos.add_argument("--output", type=Path, default=RAIZ / "data")
    salida = argumentos.parse_args().output.resolve()
    salida.mkdir(parents=True, exist_ok=True)
    base = salida / "concentraciones_mM.db"
    if base.exists():
        raise FileExistsError(
            (
                "Use una copia con data vacía para recons"
                "truir; no se sobrescribe SQLite."
            )
        )
    huellas = {
        archivo.name: huella(archivo)
        for archivo in (RAIZ / "sources").iterdir()
        if archivo.is_file()
    }
    libro_mmc2 = openpyxl.load_workbook(
        RAIZ / "sources/mmc2.xlsx", data_only=True
    )
    libro_mmc3 = openpyxl.load_workbook(
        RAIZ / "sources/mmc3.xlsx", data_only=True
    )
    libro_ishii = openpyxl.load_workbook(
        RAIZ / "sources/ishii.xlsx", data_only=True
    )
    hoja_conversion = libro_mmc2["Conversion Factors"]
    volumen = float(hoja_conversion["B6"].value)
    if not (
        volumen == 2.3
        and hoja_conversion["B5"].value == "All conditions"
        and ("ml/gCDW" in hoja_conversion["A6"].value)
    ):
        raise ValueError(
            ("Factor de volumen o unidad diferentes a la fuente verificada.")
        )
    texto_metodos = (
        PdfReader(RAIZ / "sources/Gerosa_methods.pdf").pages[8].extract_text()
    )
    if not (
        "intracellular cytosolic volume" in texto_metodos
        and "µmol/gCDW" in texto_metodos
    ):
        raise ValueError(
            ("No se encontró el respaldo del volumen en los métodos.")
        )
    factor = dict(
        factor_id="GEROSA_CYTOSOLIC_VOLUME",
        value=volumen,
        unit="mL/gCDW",
        definition="intracellular cytosolic volume per dry biomass",
        scope="All conditions; MMC2 and MMC3 intracellular metabolites",
        source_file="mmc2.xlsx",
        source_sheet="Conversion Factors",
        source_cell="B6",
        definition_cells="A5:B6",
        methods_source="Gerosa_methods.pdf",
        methods_pdf_page=9,
        methods_section="Intracellular metabolite concentrations by LC-MS/MS",
        factor_sd=None,
        factor_uncertainty_status="not_reported",
        origin_limit=(
            "Source says previously determined; origi"
            "nal measurement protocol and uncertainty"
            " not specified here"
        ),
    )
    fuente = RAIZ / "sources/ecoli_auditada_v3.db"
    with sqlite3.connect(
        fuente.resolve().as_uri() + "?mode=ro", uri=True
    ) as conexion:
        mediciones = pd.read_sql_query(
            (
                "SELECT * FROM audit_measurements_readabl"
                "e WHERE entity_type='metabolite'"
            ),
            conexion,
        )
    if not (
        len(mediciones) == 20483
        and (not mediciones.measurement_uid.duplicated().any())
    ):
        raise ValueError(("Conteo inesperado o mediciones duplicadas."))
    mediciones["source_file"] = mediciones.dataset_id.map(
        {"MMC2": "mmc2.xlsx", "MMC3": "mmc3.xlsx", "ISHII2007": "ishii.xlsx"}
    )
    mediciones["source_file_sha256"] = mediciones.source_file.map(huellas)
    libros = {"MMC2": libro_mmc2, "MMC3": libro_mmc3, "ISHII2007": libro_ishii}
    for registro in mediciones.itertuples():
        valor_fuente = libros[registro.dataset_id][registro.source_sheet][
            registro.source_cell
        ].value
        if pd.notna(registro.value_original):
            if not (
                isinstance(valor_fuente, (int, float))
                and np.isclose(
                    valor_fuente,
                    registro.value_original,
                    rtol=1e-12,
                    atol=1e-14,
                )
            ):
                raise ValueError((registro.measurement_uid, valor_fuente))
        if registro.dataset_id in ["MMC2", "MMC3"] and pd.notna(
            registro.uncertainty_original
        ):
            desviacion_fuente = libros[registro.dataset_id][
                registro.source_sheet
            ][registro.source_uncertainty_cell].value
            if not np.isclose(
                desviacion_fuente,
                registro.uncertainty_original,
                rtol=1e-12,
                atol=1e-14,
            ):
                raise ValueError(
                    ("La incertidumbre no coincide con la celda de origen.")
                )
    es_gerosa = mediciones.dataset_id.isin(["MMC2", "MMC3"])
    if set(mediciones.loc[es_gerosa, "unit_original"]) != {"umol gCDW^-1"}:
        raise ValueError(("Las unidades Gerosa no corresponden a µmol/gCDW."))
    if set(mediciones.loc[~es_gerosa, "unit_original"]) != {"mM"}:
        raise ValueError(("Las unidades Ishii no corresponden a mM."))
    mediciones["conversion_divisor"] = np.where(es_gerosa, volumen, 1.0)
    mediciones["factor_id"] = np.where(
        es_gerosa, factor["factor_id"], "IDENTITY_mM"
    )
    mediciones["intracellular_volume_mL_gCDW"] = np.where(
        es_gerosa, volumen, np.nan
    )
    mediciones["unit_final"] = "mM"
    mediciones["value_mM"] = (
        mediciones.value_original / mediciones.conversion_divisor
    )
    mediciones["uncertainty_mM"] = (
        mediciones.uncertainty_original / mediciones.conversion_divisor
    )
    mediciones["conversion_formula"] = np.where(
        es_gerosa, "value_original / 2.3", "value_original (unchanged)"
    )
    mediciones["conversion_uncertainty_note"] = np.where(
        es_gerosa,
        "SD scaled with fixed V; excludes unreported uncertainty of V",
        "Original uncertainty retained without adding estimates",
    )
    mediciones["processing_stage"] = (
        "unit_conversion_only_no_reference_normalization_no_imputation"
    )
    if not mediciones.value_mM.isna().equals(mediciones.value_original.isna()):
        raise ValueError(("La conversión alteró el patrón de faltantes."))
    if not np.allclose(
        mediciones.value_mM * mediciones.conversion_divisor,
        mediciones.value_original,
        equal_nan=True,
    ):
        raise ValueError(
            ("La conversión no recupera las concentraciones originales.")
        )
    if not np.array_equal(
        mediciones.loc[~es_gerosa, "value_mM"].fillna(-999),
        mediciones.loc[~es_gerosa, "value_original"].fillna(-999),
    ):
        raise ValueError(("Las concentraciones Ishii fueron modificadas."))
    if not (mediciones.is_imputed == 0).all():
        raise ValueError("La fuente contiene valores imputados.")
    with sqlite3.connect(base) as conexion:
        mediciones.to_sql("metabolite_measurements_mM", conexion, index=False)
        pd.DataFrame([factor]).to_sql(
            "conversion_factors", conexion, index=False
        )
        conexion.executescript(
            (
                "CREATE UNIQUE INDEX measurement_uid ON m"
                "etabolite_measurements_mM(measurement_ui"
                "d); CREATE VIEW intracellular_mM AS SELE"
                "CT * FROM metabolite_measurements_mM WHE"
                "RE compartment='intracellular'; CREATE V"
                "IEW extracellular_mM AS SELECT * FROM me"
                "tabolite_measurements_mM WHERE compartme"
                "nt='extracellular';"
            )
        )
        if conexion.execute("pragma integrity_check").fetchone()[0] != "ok":
            raise ValueError(
                ("La base convertida no supera la comprobación de integridad.")
            )
    for estudio, tabla in mediciones.groupby("dataset_id"):
        tabla.to_csv(
            salida / (estudio + "_mM.csv"),
            index=False,
            encoding="utf-8-sig",
            float_format="%.17g",
        )
    pd.DataFrame([factor]).to_csv(
        salida / "conversion_factors.csv", index=False, encoding="utf-8-sig"
    )
    (salida / "factor_evidence.txt").write_text(
        "MMC2 Conversion Factors A5:B6: "
        + str(
            [
                [hoja_conversion.cell(i, j).value for j in [1, 2]]
                for i in [5, 6]
            ]
        )
        + "\n\nGerosa methods, PDF page 9:\n"
        + texto_metodos
    )
    resumen = mediciones.groupby(
        ["dataset_id", "compartment"], as_index=False
    ).agg(cells=("measurement_uid", "size"), observed=("value_mM", "count"))
    resumen["missing"] = resumen.cells - resumen.observed
    resumen.to_csv(salida / "summary.csv", index=False)
    if huellas != {
        archivo.name: huella(archivo)
        for archivo in (RAIZ / "sources").iterdir()
        if archivo.is_file()
    }:
        raise ValueError(
            ("Los archivos fuente cambiaron durante la conversión.")
        )
    with sqlite3.connect(
        fuente.resolve().as_uri() + "?mode=ro", uri=True
    ) as conexion:
        metadatos = pd.read_sql_query(
            "SELECT * FROM ishii_samples ORDER BY sample_id", conexion
        )
    metadatos.to_csv(
        salida / "samples_ishii.csv", index=False, encoding="utf-8-sig"
    )
    manifest = dict(
        source_sha256=huellas,
        script_sha256=huella(Path(__file__)),
        factor=factor,
        checks={
            "all_original_values_checked_against_workbooks": True,
            "Gerosa_SD_checked": True,
            "missing_pattern_preserved": True,
            "Ishii_unchanged": True,
            "sources_unchanged": True,
            "reverse_conversion_passed": True,
        },
        python=sys.version,
        output_sha256={
            archivo.name: huella(archivo)
            for archivo in salida.iterdir()
            if archivo.is_file()
        },
    )
    (salida / "manifest_conversion.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False)
    )
    print(resumen.to_string(index=False))


if __name__ == "__main__":
    main()
